# Builds the three training workbooks from the JSON printed by export.js (all data invented).
#   node docs/training/export.js | python3 docs/training/build-xlsx.py docs/training
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

out = sys.argv[1] if len(sys.argv) > 1 else '.'
D = json.load(sys.stdin)
F = 'Arial'
thin = Side(style='thin', color='999999')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='D9E2F3')
NOTE = PatternFill('solid', fgColor='FFF2CC')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
DAYS = ['الإثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة']
HOURS = ['8 - 9', '9 - 10', '10 - 11', '14 - 15', '15 - 16', '16 - 17']
FAKE = 'بيانات وهمية للتدريب فقط — لا علاقة لها بأي مؤسسة أو شخص حقيقي'


def font(c, bold=False, size=11, color='000000'):
    c.font = Font(name=F, bold=bold, size=size, color=color)


def lists(path, key, title):
    wb = Workbook(); wb.remove(wb.active)
    for cl in D['classes']:
        ws = wb.create_sheet(cl['name']); ws.sheet_view.rightToLeft = True
        ws['A1'] = D['inst'] + ' — ' + title; font(ws['A1'], True, 13)
        ws['A2'] = 'القسم: ' + cl['name']; font(ws['A2'], True, 12)
        ws['A3'] = FAKE; font(ws['A3'], False, 10, '7F6000'); ws['A3'].fill = NOTE
        for c, h in enumerate(['الرقم', 'الاسم الكامل', 'تاريخ الازدياد'], 1):
            x = ws.cell(row=5, column=c, value=h); font(x, True); x.fill = HEAD; x.alignment = CENTER; x.border = BOX
        for i, s in enumerate(cl[key]):
            for c, v in enumerate([i + 1, s['n'], s['b']], 1):
                x = ws.cell(row=6 + i, column=c, value=v); font(x); x.border = BOX
                x.alignment = Alignment(horizontal='center' if c != 2 else 'right')
        ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 16
    wb.save(os.path.join(out, path))


def org(path):
    wb = Workbook()
    # 1) grid: sections x (day, hour) -> lesson code
    ws = wb.active; ws.title = 'جدول الأقسام'; ws.sheet_view.rightToLeft = True
    ws['A1'] = D['inst'] + ' — التنظيم التربوي — الموسم الدراسي 2026-2027'; font(ws['A1'], True, 14)
    ws['A2'] = FAKE; font(ws['A2'], False, 10, '7F6000'); ws['A2'].fill = NOTE
    ws.cell(row=4, column=1, value='القسم'); ws.cell(row=5, column=1, value='')
    ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    for d, name in enumerate(DAYS):
        c0 = 2 + d * 6
        ws.cell(row=4, column=c0, value=name)
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 5)
        for k, h in enumerate(HOURS):
            ws.cell(row=5, column=c0 + k, value=h)
    for r in (4, 5):
        for c in range(1, 32):
            x = ws.cell(row=r, column=c); font(x, True); x.fill = HEAD; x.alignment = CENTER; x.border = BOX
    for i, g in enumerate(D['grid']):
        r = 6 + i
        x = ws.cell(row=r, column=1, value=g['name']); font(x, True); x.border = BOX; x.alignment = CENTER
        for s, code in enumerate(g['slots']):
            x = ws.cell(row=r, column=2 + s, value=code or None); font(x, False, 10); x.alignment = CENTER; x.border = BOX
    ws.column_dimensions['A'].width = 12
    for c in range(2, 32): ws.column_dimensions[L(c)].width = 8
    ws.freeze_panes = 'B6'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    n = 6 + len(D['grid']) + 1
    ws.cell(row=n, column=1, value='ملحوظة: الرمز = أستاذ (انظروا ورقة «الأساتذة»)؛ ACT/AR1 = حصة أنشطة مع أستاذ اللغة العربية AR1.'); font(ws.cell(row=n, column=1), False, 10, '555555')

    # 2) the Arabic timetable of each section
    ws2 = wb.create_sheet('استعمال الزمن'); ws2.sheet_view.rightToLeft = True
    r = 1
    for ct in D['classTimetable']:
        ws2.cell(row=r, column=1, value='استعمال الزمن لقسم ' + ct['name']); font(ws2.cell(row=r, column=1), True, 12)
        r += 1
        x = ws2.cell(row=r, column=1, value='اليوم'); font(x, True); x.fill = HEAD; x.border = BOX; x.alignment = CENTER
        for k, h in enumerate(HOURS):
            x = ws2.cell(row=r, column=2 + k, value=h); font(x, True); x.fill = HEAD; x.border = BOX; x.alignment = CENTER
        r += 1
        for d, name in enumerate(DAYS):
            x = ws2.cell(row=r, column=1, value=name); font(x, True); x.border = BOX; x.alignment = CENTER
            for k in range(6):
                x = ws2.cell(row=r, column=2 + k, value=ct['slots'][d * 6 + k] or None); font(x, False, 10); x.border = BOX; x.alignment = CENTER
            r += 1
        r += 2
    ws2.column_dimensions['A'].width = 12
    for c in range(2, 8): ws2.column_dimensions[L(c)].width = 22

    # 3) one block per teacher (the subject is written as in the staff list: a few spellings are deliberately wrong)
    ws3 = wb.create_sheet('الأساتذة'); ws3.sheet_view.rightToLeft = True
    for i, t in enumerate(D['teachers']):
        r0 = 1 + i * 6; c0 = 1  # one block per band of rows: the importer looks for the code on the same row band
        vals = ['جدول حصص', t['name'], 'مادة : ' + t['subject'], t['code'], 'الحصص الأسبوعية: ' + str(t['hours'])]
        for k, v in enumerate(vals):
            x = ws3.cell(row=r0 + k, column=c0, value=v); font(x, k in (0, 1), 11 if k else 10); x.border = BOX; x.alignment = CENTER
            if k == 0: x.fill = HEAD
            if k == 3: font(x, True, 11, '1F4E78')
    ws3.column_dimensions['A'].width = 34
    wb.save(os.path.join(out, path))


lists('لوائح_التلاميذ_تجريبية.xlsx', 'students', 'لوائح التلاميذ')
lists('لوائح_التلاميذ_محدّثة_تجريبية.xlsx', 'updated', 'لوائح التلاميذ (بعد التحديث)')
org('التنظيم_التربوي_تجريبي.xlsx')
print('ok')
